# -*- coding: utf-8 -*-
"""Agrega la hoja PLANILLA MAESTRA al Excel de tanques. No modifica las hojas existentes."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = 'tanque.xlsx'
OUT = 'calculo_de_volumen_tanque.xlsx'

wb = openpyxl.load_workbook(SRC)
if 'PLANILLA MAESTRA' in wb.sheetnames:
    del wb['PLANILLA MAESTRA']
ws = wb.create_sheet('PLANILLA MAESTRA', 0)

# ---------- estilos ----------
F = 'Calibri'
TITULO   = Font(name=F, size=18, bold=True, color='1F3864')
SUBT     = Font(name=F, size=10, italic=True, color='595959')
H_FONT   = Font(name=F, size=10, bold=True, color='FFFFFF')
H_FILL   = PatternFill('solid', fgColor='1F3864')
CELL     = Font(name=F, size=10)
CELL_IN  = Font(name=F, size=10, color='0000FF')          # dato cargado a mano
CELL_WARN= Font(name=F, size=10, color='C00000', bold=True)
SEC      = Font(name=F, size=12, bold=True, color='1F3864')
LEG      = Font(name=F, size=9, color='404040')
TOT_FONT = Font(name=F, size=11, bold=True)
TOT_FILL = PatternFill('solid', fgColor='D9E2F3')
AMARILLO = PatternFill('solid', fgColor='FFFF00')          # completar
GRIS     = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CTR  = Alignment(horizontal='center', vertical='center')

anchos = {'A':10,'B':17,'C':18,'D':15,'E':12,'F':11,'G':11,'H':13,'I':13,
          'J':14,'K':16,'L':17,'M':16,'N':13,'O':16,'P':18,'Q':15,'R':62}
for c,w in anchos.items():
    ws.column_dimensions[c].width = w

# ---------- titulo ----------
ws['A1'] = 'PLANILLA MAESTRA DE DEPÓSITOS DE COMBUSTIBLE'
ws['A1'].font = TITULO
ws.merge_cells('A1:R1')
ws.row_dimensions[1].height = 26

ws['A2'] = ('Fuente: hojas de medición de este mismo archivo + notas manuscritas "Registro de despacho de combustible". '
            'Esta hoja es la lista única de tanques que va a leer el sistema. Las hojas de cada tanque NO se modificaron.')
ws['A2'].font = SUBT
ws.merge_cells('A2:R2')
ws['A2'].alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[2].height = 26

# ---------- leyenda ----------
ws['A4'] = 'CÓMO USAR ESTA HOJA'
ws['A4'].font = SEC
leyenda = [
 'Las celdas AMARILLAS son las que hay que completar a mano. Todo lo demás ya está cargado o se calcula solo.',
 'Los números en AZUL se escribieron a mano. Los números en NEGRO son fórmulas: si cambiás la longitud o el diámetro, la capacidad se recalcula sola.',
 'La columna CAPACIDAD usa la fórmula del cilindro acostado:  Pi x (diámetro / 2)² x longitud x 1000.  Todas las medidas en metros, el resultado en litros.',
 'La columna LTS x CM DE VARILLA dice cuántos litros representa 1 cm de varilla en el punto más ancho del tanque (la mitad). Sirve para saber cuánto error mete una lectura floja.',
 'La columna ESTADO DEL DATO avisa si esa fila está lista (OK), si hay que corregir algo (CORREGIR) o si falta información (FALTAN DATOS).',
]
r = 5
for t in leyenda:
    ws.cell(row=r, column=1, value='•  ' + t).font = LEG
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=18)
    r += 1

# ---------- encabezados tabla ----------
HDR_ROW = 11
headers = ['ID','CAMPO /\nESTABLECIMIENTO','NOMBRE DEL\nTANQUE','TIPO','COMBUSTIBLE',
           'LONGITUD\n(m)','DIÁMETRO\n(m)','CAPACIDAD\n(Lts)','LTS x CM\nDE VARILLA',
           '¿TIENE\nCONTADOR?','TIPO DE\nCONTADOR','¿TIENE VARILLA /\nTABLA DE AFORO?',
           'HOJA DE AFORO\n(en este archivo)','UNIDAD DE\nLA VARILLA',
           'CÓDIGO DE ACCESO\n(PROPUESTO)','DESPACHANTE\n(nombre)','ESTADO\nDEL DATO','OBSERVACIONES']
for i,h in enumerate(headers, start=1):
    c = ws.cell(row=HDR_ROW, column=i, value=h)
    c.font = H_FONT; c.fill = H_FILL; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c.border = BOX
ws.row_dimensions[HDR_ROW].height = 42

# ---------- datos ----------
# (campo, tanque, tipo, comb, L, D, capacidad_fija, contador, tipo_cont, aforo, hoja, unidad, codigo, estado, obs)
datos = [
 ('LA UNION','LA UNION','A VERIFICAR','GASOIL',4.78,1.74,None,'A VERIFICAR','','SÍ','LA UNION','cm','UNION26','VERIFICAR',
  'Campo que NO tiene código en la app de balanza: hay que crear uno nuevo. OJO: la tabla de aforo de este tanque está en CENTÍMETROS (0 a 174), mientras que todas las demás están en METROS. Conviene pasarlas todas a la misma unidad.'),
 ('LA PURIFICADA','TANQUE GRANDE','FIJO','GASOIL',5.5,2.28,None,'SÍ (a confirmar)','MECÁNICO','SÍ','LA PURIF','m','PURIFICADA26','OK',
  'Antes se llamaba QUIMILI. Renombrado a LA PURIFICADA según la nota manuscrita.'),
 ('LA PURIFICADA','TANQUE CHICO','FIJO','GASOIL',3.15,1.48,None,'SÍ (a confirmar)','MECÁNICO','SÍ','LA PURIF chico','m','PURIFICADA26','OK',
  'Segundo tanque del mismo campo. El sistema tiene que poder distinguirlo del grande al momento de despachar.'),
 ('EL BUFALO','EL BUFALO','FIJO','GASOIL',7.62,2.28,None,'SÍ (a confirmar)','MECÁNICO','SÍ','EL BUFALO','m','BUFALO26','CORREGIR',
  'Campo que NO tiene código en la app de balanza: hay que crear uno nuevo. ERROR DE TIPEO EN LA TABLA DE AFORO: a 0,15 m la hoja dice 589 Lts y le corresponden 873 Lts (falta 284 Lts). Es el tanque más grande con aforo, y el que más litros mueve por cm de varilla.'),
 ('EL C1','CISTERNA 1','A VERIFICAR','GASOIL',4.95,2.25,None,'A VERIFICAR','','SÍ','EL C1','m','C126','VERIFICAR',
  'La hoja lo llama "Cisterna", pero la nota manuscrita dice que en el Excel están los depósitos FIJOS y que las cisternas faltan listar. Hay que definir si este tanque es fijo o móvil: si es móvil, se traslada entre campos y eso cambia cómo se registra.'),
 ('EL C1','CISTERNA 2','A VERIFICAR','GASOIL',5.05,1.96,None,'A VERIFICAR','','SÍ','EL C1 (2)','m','C126','CORREGIR',
  'Ídem anterior (definir fijo o móvil). ERROR DE TIPEO EN LA TABLA DE AFORO: a 0,05 m la hoja dice 50 Lts y le corresponden 105 Lts. ADEMÁS: la hoja "EL C1 (3) Tanquesito" es una COPIA de este tanque (misma longitud 5,05 y mismo diámetro 1,96, misma tabla entera). Confirmado con vos que EL C1 tiene 2 depósitos, no 3: esa hoja se puede borrar o hay que medir el tercero de verdad.'),
 ('EL WICHI','WICHI','FIJO','GASOIL',6.08,2.5,None,'SÍ (a confirmar)','MECÁNICO','SÍ','WICHI','m','WICHI26','OK',
  'Medidas y tabla de aforo verificadas: coinciden con la fórmula del cilindro.'),
 ('LA PRADERA','TANQUE 20 MIL','FIJO','GASOIL',4.5,2.4,None,'SÍ (a confirmar)','MECÁNICO','SÍ','LA PRADERA','m','PRADERA26','OK',
  'Uno de los dos tanques de LA PRADERA. Medidas y tabla de aforo verificadas.'),
 ('LA PRADERA','TANQUE 60 MIL','FIJO','GASOIL',None,None,60000,'A VERIFICAR','','NO','FALTA','','PRADERA26','FALTAN DATOS',
  'Segundo tanque de LA PRADERA. NO tiene hoja de medición en este archivo: hay que medir longitud y diámetro para calcular la capacidad real y armar su tabla de aforo. Los 60.000 Lts son un estimado: la hoja "total" anotaba 50.000 y 60.000 con la nota "se creía de este tamaño, por medidas daba 60 mil".'),
]

row = HDR_ROW + 1
first_data = row
for i,(campo,tanque,tipo,comb,L,D,cap_fija,cont,tcont,aforo,hoja,unid,cod,estado,obs) in enumerate(datos, start=1):
    ws.cell(row=row, column=1, value=f'COMB-{i:02d}')
    ws.cell(row=row, column=2, value=campo)
    ws.cell(row=row, column=3, value=tanque)
    ws.cell(row=row, column=4, value=tipo)
    ws.cell(row=row, column=5, value=comb)

    cL = ws.cell(row=row, column=6, value=L)
    cD = ws.cell(row=row, column=7, value=D)
    for c in (cL, cD):
        c.font = CELL_IN; c.number_format = '0.00'
        if c.value is None:
            c.fill = AMARILLO

    cH = ws.cell(row=row, column=8)
    if cap_fija is None:
        cH.value = f'=ROUND(PI()*(G{row}/2)^2*F{row}*1000,0)'
        cH.font = CELL
    else:
        cH.value = cap_fija
        cH.font = CELL_IN
    cH.number_format = '#,##0'

    cI = ws.cell(row=row, column=9)
    if cap_fija is None:
        cI.value = f'=ROUND(G{row}*F{row}*10,0)'
        cI.font = CELL
    else:
        cI.value = None
        cI.fill = AMARILLO
    cI.number_format = '#,##0'

    ws.cell(row=row, column=10, value=cont)
    ws.cell(row=row, column=11, value=tcont)
    ws.cell(row=row, column=12, value=aforo)
    ws.cell(row=row, column=13, value=hoja)
    ws.cell(row=row, column=14, value=unid)
    ws.cell(row=row, column=15, value=cod)
    ws.cell(row=row, column=16, value=None).fill = AMARILLO   # despachante
    ws.cell(row=row, column=17, value=estado)
    ws.cell(row=row, column=18, value=obs)

    for col in range(1,19):
        c = ws.cell(row=row, column=col)
        c.border = BOX
        if c.font is None or c.font.color is None or c.font.color.rgb not in ('FF0000FF',):
            if col not in (6,7,8,9):
                c.font = CELL
        if col == 18:
            c.alignment = WRAP
        elif col in (1,4,5,6,7,8,9,10,12,14,15,17):
            c.alignment = CTR
        else:
            c.alignment = Alignment(vertical='center')
    if estado in ('CORREGIR','FALTAN DATOS'):
        ws.cell(row=row, column=17).font = CELL_WARN
    ws.row_dimensions[row].height = 58
    row += 1
last_data = row - 1

# ---------- total ----------
ws.cell(row=row, column=3, value='TOTAL CAPACIDAD INSTALADA').font = TOT_FONT
tot = ws.cell(row=row, column=8, value=f'=SUM(H{first_data}:H{last_data})')
tot.font = TOT_FONT; tot.number_format = '#,##0'
ws.cell(row=row, column=18, value='Suma de los 9 tanques listados, tomando los 60.000 Lts estimados de LA PRADERA. La hoja "total" del archivo original da 215.474 Lts: '
        'la diferencia de 1 litro es porque ahí el tanque chico de LA PURIFICADA figura redondeado a 5.420 y el cálculo exacto da 5.419,3 Lts.').font = LEG
ws.cell(row=row, column=18).alignment = WRAP
for col in range(1,19):
    c = ws.cell(row=row, column=col); c.fill = TOT_FILL; c.border = BOX
total_row = row
row += 2

# ---------- cisternas a completar ----------
ws.cell(row=row, column=1, value='CISTERNAS MÓVILES POR LISTAR — completar las filas amarillas').font = SEC
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=18)
row += 1
ws.cell(row=row, column=1, value='La nota manuscrita dice: "Hay otros depósitos móviles, que llamamos cisternas, en varios campos. A veces se trasladan entre ellos" y "Falta listar las cisternas". '
        'Cada cisterna necesita: en qué campo está hoy, capacidad, y si tiene contador. Como se trasladan, el sistema las va a tratar como depósitos que cambian de campo (movimiento tipo TRASLADO).').font = LEG
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=18)
ws.cell(row=row, column=1).alignment = WRAP
ws.row_dimensions[row].height = 30
row += 1
for i in range(10, 16):
    ws.cell(row=row, column=1, value=f'COMB-{i:02d}').font = CELL
    for col in range(2,19):
        c = ws.cell(row=row, column=col)
        c.fill = AMARILLO; c.border = BOX; c.font = CELL
    ws.cell(row=row, column=1).border = BOX
    ws.cell(row=row, column=4, value='CISTERNA MÓVIL').font = CELL
    ws.cell(row=row, column=5, value='GASOIL').font = CELL
    ws.cell(row=row, column=8, value=f'=IF(AND(F{row}<>"",G{row}<>""),ROUND(PI()*(G{row}/2)^2*F{row}*1000,0),"")')
    ws.cell(row=row, column=8).number_format = '#,##0'
    ws.cell(row=row, column=9, value=f'=IF(AND(F{row}<>"",G{row}<>""),ROUND(G{row}*F{row}*10,0),"")')
    ws.cell(row=row, column=9).number_format = '#,##0'
    row += 1
row += 1

# ---------- revision ----------
ws.cell(row=row, column=1, value='REVISIÓN DE LA MEDICIÓN Y EL CÁLCULO DE VOLUMEN').font = SEC
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=18)
row += 1

revision = [
 ('QUÉ SE VERIFICÓ',
  'Se recalcularon los 9 tanques con la fórmula del cilindro acostado y se compararon TODOS los puntos de las 9 tablas de aforo '
  '(0 a diámetro completo) contra la fórmula exacta del volumen parcial de un cilindro horizontal.'),
 ('LA CAPACIDAD ESTÁ BIEN',
  'Los 9 cálculos de capacidad son correctos: la fórmula Pi x radio² x longitud x 1000 está bien aplicada en todas las hojas. '
  'No hay que corregir ninguna capacidad (salvo el tanque de 60 mil de LA PRADERA, que nunca se midió).'),
 ('LAS TABLAS DE AFORO ESTÁN BIEN, CON 2 ERRORES DE TIPEO',
  'De casi 400 puntos revisados, solo 2 están mal: (1) EL BUFALO a 0,15 m dice 589 Lts y corresponden 873 Lts; '
  '(2) EL C1 (2) a 0,05 m dice 50 Lts y corresponden 105 Lts. Todo el resto coincide con la fórmula (diferencias de 1 a 5 litros por redondeo).'),
 ('PROBLEMA DE UNIDADES',
  'El encabezado de todas las hojas dice "Cm de varilla", pero solo LA UNION está realmente en centímetros (0 a 174). '
  'Las otras 8 están en METROS (0 a 2,28 / 2,50 / etc.). Es un error de rótulo, no de cálculo, pero si alguien lee la tabla '
  'creyendo que son centímetros se equivoca por 100 veces. Hay que unificar a centímetros antes de cargarlo al sistema.'),
 ('HOJA DUPLICADA',
  'La hoja "EL C1 (3) Tanquesito" tiene exactamente las mismas medidas y la misma tabla que "EL C1 (2)". '
  'Confirmado que EL C1 tiene 2 depósitos: o se borra esa hoja, o se mide el tercer tanque de verdad.'),
 ('LO QUE LA FÓRMULA NO CONTEMPLA',
  'La fórmula trata al tanque como un cilindro perfecto. En la realidad: (a) los cabezales curvos (abombados) suman litros que no están contados, '
  'normalmente entre 1% y 3% del total; (b) si el tanque no está perfectamente nivelado, la varilla lee distinto según de qué punta se mida; '
  '(c) el diámetro cargado parece ser el EXTERIOR, y el volumen real se calcula con el interior (la chapa resta unos litros). '
  'Para control de stock esto es aceptable, pero explica por qué la varilla y el cálculo nunca van a coincidir exactamente.'),
 ('CUÁNTO ERROR METE LA VARILLA',
  'La varilla está marcada cada 5 cm. En el punto más ancho del tanque, 1 cm de error de lectura vale: 83 Lts en LA UNION, '
  '174 Lts en EL BUFALO, 152 Lts en EL WICHI. O sea que un salto de 5 cm en EL BUFALO son casi 870 litros. '
  'Es la razón por la que la varilla sirve para controlar, no para medir un despacho.'),
 ('LO QUE FALTA PARA CERRAR ESTA PLANILLA',
  '1) Medir el tanque de 60 mil de LA PRADERA. 2) Listar todas las cisternas móviles. 3) Definir, tanque por tanque, si tiene contador y si funciona. '
  '4) Poner el nombre del despachante de cada depósito. 5) Confirmar los códigos de acceso.'),
 ('LOS CAMPOS NO COINCIDEN CON LOS DE LA BALANZA',
  'La app de balanza tiene código para: GENERAL, EL MATACO, LA PRADERA, EL C1, EL WICHI, LA JUANITA, QUIMILI y PACO-PASCUAL. '
  'Pero los tanques de este Excel están en: LA UNION, LA PURIFICADA, EL BUFALO, EL C1, EL WICHI y LA PRADERA. '
  'Consecuencia: LA UNION y EL BUFALO necesitan código nuevo (no existen en balanza), y EL MATACO y LA JUANITA tienen código pero no tienen tanque listado '
  '(¿no tienen depósito, o falta cargarlo?). QUIMILI pasa a llamarse LA PURIFICADA y PACO-PASCUAL se elimina, según la nota manuscrita.'),
 ('CÓDIGOS DE ACCESO PROPUESTOS',
  'Según la nota ("nombres del campo más el número 26"), la propuesta es: UNION26, PURIFICADA26, BUFALO26, C126, WICHI26, PRADERA26, '
  'más MATACO26 y JUANITA26 si esos campos tienen depósito, y GENERAL26 para el usuario general. '
  'Falta definir el segundo código de cada campo, el de "solo ver registros", que no puede cargar nada.'),
]
for titulo, texto in revision:
    c1 = ws.cell(row=row, column=1, value=titulo)
    c1.font = Font(name=F, size=10, bold=True); c1.alignment = WRAP; c1.fill = GRIS; c1.border = BOX
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    for col in range(1,5):
        ws.cell(row=row, column=col).fill = GRIS; ws.cell(row=row, column=col).border = BOX
    c2 = ws.cell(row=row, column=5, value=texto)
    c2.font = CELL; c2.alignment = WRAP; c2.border = BOX
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=18)
    for col in range(5,19):
        ws.cell(row=row, column=col).border = BOX
    ws.row_dimensions[row].height = 46
    row += 1

ws.freeze_panes = f'A{HDR_ROW+1}'
ws.sheet_view.showGridLines = False
wb.save(OUT)
print('OK ->', OUT, '| filas datos', first_data, '-', last_data, '| total en fila', total_row)
